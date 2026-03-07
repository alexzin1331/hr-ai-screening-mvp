import os
import json
import uuid
import zipfile
import asyncio
import logging
import re
from dotenv import load_dotenv
from fastapi import FastAPI, UploadFile, HTTPException
from pdfminer.high_level import extract_text
from docx import Document
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from groq import Groq
from telegram import Bot

# Загружаем переменные из .env файла
load_dotenv()

# ------------------------
# LOGGING
# ------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
)
logger = logging.getLogger("hr-ai-screening")

# ------------------------
# CONFIG
# ------------------------
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
BOT_TOKEN = os.getenv("BOT_TOKEN")

if not GROQ_API_KEY:
    logger.error("GROQ_API_KEY не найден в переменных окружения")
    raise ValueError("GROQ_API_KEY не найден в переменных окружения")
if not BOT_TOKEN:
    logger.error("BOT_TOKEN не найден в переменных окружения")
    raise ValueError("BOT_TOKEN не найден в переменных окружения")

logger.info("Environment variables loaded successfully. Initializing Groq and Telegram clients.")

# Инициализируем клиенты
client = Groq(api_key=GROQ_API_KEY)
bot = Bot(token=BOT_TOKEN)

app = FastAPI()

# Отдаём статику дашборда
app.mount("/dashboard", StaticFiles(directory="dashboard"), name="dashboard")


@app.get("/")
async def root():
    return FileResponse("dashboard/index.html")


@app.get("/app.js")
async def app_js():
    return FileResponse("dashboard/app.js", media_type="application/javascript")


# ------------------------
# MEMORY STORAGE
# ------------------------
vacancy = {}
candidates = {}


# ------------------------
# PARSERS
# ------------------------
def parse_pdf(path):
    logger.info("parse_pdf: starting PDF text extraction for file: %s", path)
    text = extract_text(path)
    logger.info("parse_pdf: completed for file: %s | length_chars=%d", path, len(text or ""))
    return text


def parse_docx(path):
    logger.info("parse_docx: starting DOCX text extraction for file: %s", path)
    doc = Document(path)
    text = "\n".join([p.text for p in doc.paragraphs])
    logger.info("parse_docx: completed for file: %s | paragraphs=%d | length_chars=%d", path, len(doc.paragraphs),
                len(text or ""))
    return text


def parse_resume(path):
    logger.info("parse_resume: dispatching parser for file: %s", path)
    try:
        if path.endswith(".pdf"):
            text = parse_pdf(path)
        elif path.endswith(".docx"):
            text = parse_docx(path)
        else:
            logger.warning("parse_resume: unsupported resume format: %s", path)
            return ""
        logger.info("parse_resume: finished for file: %s | length_chars=%d", path, len(text or ""))
        return text
    except Exception:
        logger.exception("parse_resume: error while parsing resume file: %s", path)
        raise HTTPException(status_code=500, detail=f"Failed to parse resume file: {os.path.basename(path)}")


# ------------------------
# TELEGRAM USERNAME EXTRACTION
# ------------------------
TG_USERNAME_PATTERN = re.compile(r"(?:https?://t\.me/|@)([A-Za-z0-9_]{5,})")


def extract_telegram_username_from_text(text: str):
    match = TG_USERNAME_PATTERN.search(text or "")
    if match:
        username = match.group(1)
        logger.info("Extracted Telegram username from raw resume text: @%s", username)
        return username
    logger.debug("No Telegram username found in raw resume text.")
    return None


# ------------------------
# HELPER: Clean LLM JSON response
# ------------------------
def clean_llm_json_response(text: str) -> str:
    """Очищает ответ LLM от маркдауна и лишнего текста"""
    text = text.strip()

    # Убираем ```json или ``` в начале/конце
    if text.startswith("```"):
        text = re.sub(r'^```(?:json)?\s*', '', text)
        text = re.sub(r'\s*```$', '', text)
        text = text.strip()

    # Если не начинается с {, ищем JSON внутри текста
    if not text.startswith("{"):
        json_match = re.search(r'\{[\s\S]*\}', text)
        if json_match:
            text = json_match.group(0)

    return text


# ------------------------
# LLM EXTRACTION
# ------------------------
def extract_candidate(resume_text):
    logger.info("Starting candidate extraction via LLM based on resume text.")

    prompt = f"""
You are an expert HR resume parser.
Extract detailed structured information from the resume.

Resume:
{resume_text}

Return STRICT JSON with these keys:
name, email, phone, location, telegram_username,
total_years_experience, current_position, seniority_level,
skills_technical, skills_soft,
programming_languages, frameworks, databases, cloud, tools,
education, languages, previous_companies, projects,
english_level, key_strengths, possible_weaknesses,
salary_expectation, summary

Return ONLY valid JSON, no explanations.
"""

    try:
        response = client.chat.completions.create(
            model="llama-3.3-70b-versatile",  # ✅ Актуальная модель
            messages=[{"role": "user", "content": prompt}],
            temperature=0,
        )
    except Exception:
        logger.exception("LLM error during candidate extraction.")
        raise HTTPException(status_code=500, detail="LLM error during candidate extraction.")

    text = response.choices[0].message.content
    logger.debug("Raw LLM candidate JSON response: %s", text)

    # 🔹 Очистка ответа
    text = clean_llm_json_response(text)

    if not text:
        logger.error("LLM returned empty response for candidate extraction")
        raise HTTPException(status_code=500, detail="LLM returned empty response")

    try:
        candidate = json.loads(text)
    except json.JSONDecodeError as e:
        logger.error("Failed to parse LLM JSON. Raw text preview: %s", text[:500])
        logger.error("JSON error: %s", str(e))
        raise HTTPException(status_code=500, detail=f"Failed to parse candidate JSON: {str(e)}")

    # Fallback для telegram_username
    if not candidate.get("telegram_username"):
        fallback_username = extract_telegram_username_from_text(resume_text)
        if fallback_username:
            candidate["telegram_username"] = fallback_username

    logger.info(
        "Candidate extracted: name=%s, email=%s, seniority=%s, total_years_experience=%s, english_level=%s, salary_expectation=%s, telegram=%s",
        candidate.get("name"), candidate.get("email"), candidate.get("seniority_level"),
        candidate.get("total_years_experience"), candidate.get("english_level"),
        candidate.get("salary_expectation"), candidate.get("telegram_username"),
    )

    # Сводка по скиллам
    tech_skills = candidate.get("skills_technical") or []
    soft_skills = candidate.get("skills_soft") or []
    tech_skills_count = len(tech_skills.split(",")) if isinstance(tech_skills, str) else len(tech_skills)
    soft_skills_count = len(soft_skills.split(",")) if isinstance(soft_skills, str) else len(soft_skills)

    logger.info("Candidate skills summary: technical_count=%d, soft_count=%d", tech_skills_count, soft_skills_count)

    return candidate


# ------------------------
# MATCHING
# ------------------------
def score_candidate(candidate):
    if not vacancy:
        logger.error("Attempt to score candidate before vacancy is created.")
        raise HTTPException(status_code=400, detail="Vacancy is not set. Create vacancy first.")

    logger.info("Scoring candidate for vacancy. vacancy_title=%s, candidate_name=%s",
                vacancy.get("title"), candidate.get("name") if isinstance(candidate, dict) else None)

    prompt = f"""
You are an HR AI that evaluates candidate fit for a vacancy.

VACANCY:
title: {vacancy['title']}
required_skills: {vacancy['skills']}
description: {vacancy['description']}
seniority: {vacancy['seniority']}

CANDIDATE:
{json.dumps(candidate, ensure_ascii=False, indent=2)}

Evaluate candidate fit. Consider: skills match, experience, seniority, relevant technologies, english, career stability.

Return STRICT JSON:
{{
  "score": <integer 0-100>,
  "match_strengths": ["strength1", "strength2"],
  "match_gaps": ["gap1", "gap2"],
  "reason": "brief explanation"
}}

Return ONLY valid JSON, no explanations.
"""

    try:
        response = client.chat.completions.create(
            model="llama-3.3-70b-versatile",  # ✅ Актуальная модель
            messages=[{"role": "user", "content": prompt}],
            temperature=0,
        )
    except Exception:
        logger.exception("LLM error during candidate scoring.")
        raise HTTPException(status_code=500, detail="LLM error during candidate scoring.")

    text = response.choices[0].message.content
    logger.debug("Raw LLM scoring JSON response: %s", text)

    # 🔹 Очистка ответа
    text = clean_llm_json_response(text)

    if not text:
        logger.error("LLM returned empty response for scoring")
        raise HTTPException(status_code=500, detail="LLM returned empty response for scoring")

    try:
        result = json.loads(text)
    except json.JSONDecodeError as e:
        logger.error("Failed to parse scoring JSON. Raw text preview: %s", text[:500])
        logger.error("JSON error: %s", str(e))
        raise HTTPException(status_code=500, detail=f"Failed to parse scoring JSON: {str(e)}")

    strengths = result.get("match_strengths")
    gaps = result.get("match_gaps")
    reason = result.get("reason")

    strengths_len = len(strengths) if isinstance(strengths, (list, str)) else 0
    gaps_len = len(gaps) if isinstance(gaps, (list, str)) else 0
    reason_preview = (reason[:200] + "...") if isinstance(reason, str) and len(reason) > 200 else reason

    logger.info("Candidate scored: score=%s, strengths_len=%d, gaps_len=%d",
                result.get("score"), strengths_len, gaps_len)
    logger.info("Candidate score reason (preview): %s", reason_preview)

    return result


# ------------------------
# TELEGRAM
# ------------------------
async def send_message(username):
    try:
        message = """Hello!

Your resume matched our vacancy.
We would like to ask you a few quick screening questions.
Please reply to start the interview."""
        # Убираем лишние @ в начале и добавляем один
        clean_username = username.lstrip('@')
        await bot.send_message(chat_id=f"@{clean_username}", text=message)
        logger.info("Telegram message successfully sent to @%s", username)
    except Exception as e:
        logger.exception("Telegram error while sending message to @%s", username)


# ------------------------
# API
# ------------------------
@app.post("/create_vacancy")
async def create_vacancy(data: dict):
    logger.info("Received create_vacancy request with data: %s", data)

    required_fields = ["title", "skills", "description", "seniority"]
    missing_fields = [f for f in required_fields if f not in data]
    if missing_fields:
        logger.error("Missing fields in create_vacancy request: %s", missing_fields)
        raise HTTPException(status_code=400, detail=f"Missing fields: {missing_fields}")

    vacancy["title"] = data["title"]
    vacancy["skills"] = data["skills"]
    vacancy["description"] = data["description"]
    vacancy["seniority"] = data["seniority"]

    logger.info("Vacancy created: title=%s, seniority=%s", vacancy["title"], vacancy["seniority"])
    return {"status": "vacancy_created"}


@app.post("/upload_resumes")
async def upload_resumes(file: UploadFile):
    logger.info("Received upload_resumes request. Filename=%s", file.filename)

    if not vacancy:
        logger.error("upload_resumes called before vacancy was created.")
        raise HTTPException(status_code=400, detail="Vacancy is not set. Create vacancy first.")

    os.makedirs("resumes", exist_ok=True)

    try:
        file_bytes = await file.read()
        with open("resumes.zip", "wb") as f:
            f.write(file_bytes)
        logger.info("Saved resumes.zip (%d bytes). Extracting...", len(file_bytes))
    except Exception:
        logger.exception("Failed to save uploaded resumes ZIP.")
        raise HTTPException(status_code=500, detail="Failed to save uploaded resumes ZIP.")

    try:
        with zipfile.ZipFile("resumes.zip", "r") as zip_ref:
            zip_ref.extractall("resumes")
        logger.info("ZIP extracted into 'resumes' directory.")
    except Exception:
        logger.exception("Failed to extract resumes ZIP.")
        raise HTTPException(status_code=500, detail="Failed to extract resumes ZIP.")

    candidates.clear()
    logger.info("Cleared previous candidates. Starting parsing pipeline.")

    processed_count = 0
    skipped_unsupported = 0
    skipped_empty = 0

    for filename in os.listdir("resumes"):
        path = f"resumes/{filename}"

        if not (path.endswith(".pdf") or path.endswith(".docx")):
            logger.warning("Skipping non-resume file inside ZIP: %s", filename)
            skipped_unsupported += 1
            continue

        logger.info("Processing resume file from ZIP: %s", filename)
        text = parse_resume(path)

        if not text:
            logger.warning("Empty or unsupported resume content: %s", filename)
            skipped_empty += 1
            continue

        candidate = extract_candidate(text)
        match = score_candidate(candidate)

        cid = str(uuid.uuid4())
        candidates[cid] = {
            "data": candidate,
            "score": match.get("score"),
            "analysis": match
        }
        processed_count += 1

        logger.info("Candidate stored. id=%s, name=%s, score=%s, telegram=%s",
                    cid, candidate.get("name"), match.get("score"), candidate.get("telegram_username"))

    logger.info("Upload and parsing completed. Total candidates parsed: %d, skipped_unsupported=%d, skipped_empty=%d",
                len(candidates), skipped_unsupported, skipped_empty)

    return {
        "parsed": len(candidates),
        "processed": processed_count,
        "skipped_unsupported": skipped_unsupported,
        "skipped_empty": skipped_empty
    }


@app.post("/start_screening")
async def start_screening(data: dict):
    logger.info("Received start_screening request with data: %s", data)

    if not vacancy:
        logger.error("start_screening called before vacancy was created.")
        raise HTTPException(status_code=400, detail="Vacancy is not set. Create vacancy first.")

    if not candidates:
        logger.error("start_screening called but no candidates are loaded.")
        raise HTTPException(status_code=400, detail="No candidates loaded. Upload resumes first.")

    try:
        n = int(data["n"])
    except (KeyError, ValueError, TypeError):
        logger.error("Invalid or missing 'n' in start_screening request: %s", data)
        raise HTTPException(status_code=400, detail="'n' must be a positive integer.")

    if n <= 0:
        logger.error("Non-positive 'n' value in start_screening: %s", n)
        raise HTTPException(status_code=400, detail="'n' must be a positive integer.")

    sorted_candidates = sorted(candidates.items(), key=lambda x: x[1]["score"], reverse=True)
    top_candidates = sorted_candidates[:n]

    logger.info("Selected top %d candidates for Telegram outreach (from total %d).",
                len(top_candidates), len(sorted_candidates))

    sent = []
    skipped_no_telegram = []

    for cid, cand in top_candidates:
        username = cand["data"].get("telegram_username")
        if not username:
            logger.warning("Candidate %s has no Telegram username, skipping.", cid)
            skipped_no_telegram.append(cid)
            continue
        await send_message(username)
        sent.append(username)

    logger.info("Screening completed. Total contacted candidates: %d, skipped_no_telegram=%d",
                len(sent), len(skipped_no_telegram))

    return {
        "contacted_candidates": sent,
        "requested_top_n": n,
        "total_candidates": len(sorted_candidates),
        "skipped_no_telegram": skipped_no_telegram,
    }


@app.get("/candidates")
def get_candidates():
    logger.info("Received get_candidates request.")
    sorted_candidates = sorted(candidates.values(), key=lambda x: x["score"], reverse=True)
    logger.info("Returning %d candidates.", len(sorted_candidates))
    return sorted_candidates


@app.get("/health")
def health():
    logger.info("Healthcheck requested.")
    return {
        "status": "ok",
        "vacancy_created": bool(vacancy),
        "candidates_loaded": len(candidates),
    }