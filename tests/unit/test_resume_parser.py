from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.core.exceptions import ResumeParseError
from app.services.resume_parser import ResumeParserService


def test_parse_docx_success(tmp_path: Path):
    file_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Python developer")
    document.save(file_path)

    parser = ResumeParserService()
    text = parser.parse_docx(file_path)

    assert "Python developer" in text


def test_parse_resume_empty_docx(tmp_path: Path):
    file_path = tmp_path / "resume.docx"
    document = Document()
    document.save(file_path)

    parser = ResumeParserService()
    with pytest.raises(ResumeParseError):
        parser.parse_resume(file_path)
