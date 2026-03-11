from __future__ import annotations

import os
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker


os.environ["DATABASE_URL"] = "sqlite:///./test.db"
os.environ["LLM_PROVIDER"] = "mock"

from app.core.config import get_settings
from app.db.session import Base, get_db
from app.main import app


@pytest.fixture()
def client(tmp_path, monkeypatch):
    settings = get_settings()
    settings.uploads_dir = tmp_path / "uploads"
    settings.uploads_dir.mkdir(parents=True, exist_ok=True)
    settings.database_url = f"sqlite:///{tmp_path / 'test.db'}"
    settings.llm_provider = "mock"
    settings.telegram_bot_token = None
    settings.telegram_bot_username = "test_bot"
    settings.email_provider = "mock"
    settings.resend_api_key = "test-key"
    settings.resend_email_from = "HR AI Screening <onboarding@resend.dev>"

    engine = create_engine(settings.database_url, connect_args={"check_same_thread": False})
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db

    import app.db.session as db_session_module
    import app.services.pipeline as pipeline_module
    import app.api.routes.health as health_module

    monkeypatch.setattr(db_session_module, "SessionLocal", TestingSessionLocal)
    monkeypatch.setattr(pipeline_module, "SessionLocal", TestingSessionLocal)
    monkeypatch.setattr(health_module, "SessionLocal", TestingSessionLocal)

    with TestClient(app) as test_client:
        yield test_client

    app.dependency_overrides.clear()


@pytest.fixture()
def sample_docx_bytes(tmp_path) -> bytes:
    from docx import Document

    doc_path = tmp_path / "candidate.docx"
    document = Document()
    document.add_paragraph("Name: Ivan Ivanov")
    document.add_paragraph("Email: ivan@example.com")
    document.add_paragraph("Phone: +79991234567")
    document.add_paragraph("Telegram: @ivan_candidate")
    document.add_paragraph("Python FastAPI PostgreSQL Docker")
    document.save(doc_path)
    return doc_path.read_bytes()


@pytest.fixture()
def empty_docx_bytes(tmp_path) -> bytes:
    from docx import Document

    doc_path = tmp_path / "empty.docx"
    document = Document()
    document.save(doc_path)
    return doc_path.read_bytes()
